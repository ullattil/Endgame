"""Renders tailored resume + cover letter data into .docx files, laid out to fit one page."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def _set_bottom_border(paragraph) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "808080")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _tight(paragraph, space_after: int = 0):
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.0
    return paragraph


def _section_heading(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    _set_bottom_border(p)
    return p


def _entry_header(doc, left_text: str, right_text: str, usable_width):
    p = doc.add_paragraph()
    _tight(p)
    p.paragraph_format.tab_stops.add_tab_stop(usable_width, WD_TAB_ALIGNMENT.RIGHT)
    run = p.add_run(left_text)
    run.bold = True
    run.font.size = Pt(10)
    if right_text:
        tab_run = p.add_run(f"\t{right_text}")
        tab_run.italic = True
        tab_run.font.size = Pt(9.5)
    return p


def _paragraph(doc, text: str, size: float = 9.5):
    p = doc.add_paragraph(text)
    _tight(p)
    for run in p.runs:
        run.font.size = Pt(size)
    return p


def _bullet(doc, text: str):
    p = doc.add_paragraph(text, style="List Bullet")
    _tight(p)
    p.paragraph_format.left_indent = Inches(0.2)
    for run in p.runs:
        run.font.size = Pt(9.5)
    return p


def render_resume(resume: dict, out_path: Path) -> None:
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    usable_width = section.page_width - section.left_margin - section.right_margin

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.0

    contact = resume.get("contact", {})

    name_p = doc.add_paragraph()
    _tight(name_p, space_after=1)
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_p.add_run(contact.get("name", ""))
    name_run.bold = True
    name_run.font.size = Pt(18)

    contact_bits = list(filter(None, [contact.get("email"), contact.get("phone"), contact.get("location")]))
    links = contact.get("links") or []
    contact_line = "  |  ".join(contact_bits + list(links))
    if contact_line:
        cp = doc.add_paragraph()
        _tight(cp)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cp.add_run(contact_line)
        run.font.size = Pt(9)

    if resume.get("summary"):
        _section_heading(doc, "Summary")
        _paragraph(doc, resume["summary"])

    if resume.get("experience"):
        _section_heading(doc, "Experience")
        for job in resume["experience"]:
            left = f"{job.get('title', '')} — {job.get('company', '')}"
            dates = f"{job.get('start', '')} – {job.get('end', '')}"
            _entry_header(doc, left, dates, usable_width)
            for bullet in job.get("bullets", []):
                _bullet(doc, bullet)

    if resume.get("projects"):
        _section_heading(doc, "Projects")
        for project in resume["projects"]:
            dates = f"{project.get('start', '')} – {project.get('end', '')}".strip(" –")
            _entry_header(doc, project.get("name", ""), dates, usable_width)
            for bullet in project.get("bullets", []):
                _bullet(doc, bullet)

    if resume.get("leadership"):
        _section_heading(doc, "Leadership Experience")
        for entry in resume["leadership"]:
            left = f"{entry.get('role', '')} — {entry.get('org', '')}"
            dates = f"{entry.get('start', '')} – {entry.get('end', '')}"
            _entry_header(doc, left, dates, usable_width)
            for bullet in entry.get("bullets", []):
                _bullet(doc, bullet)

    if resume.get("education"):
        _section_heading(doc, "Education")
        for edu in resume["education"]:
            _paragraph(doc, f"{edu.get('degree', '')}, {edu.get('school', '')} ({edu.get('year', '')})")

    if resume.get("skills"):
        _section_heading(doc, "Skills")
        _paragraph(doc, ", ".join(resume["skills"]))

    if resume.get("certifications"):
        _section_heading(doc, "Certifications")
        _paragraph(doc, ", ".join(resume["certifications"]))

    if resume.get("awards"):
        _section_heading(doc, "Awards")
        _paragraph(doc, ", ".join(resume["awards"]))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def render_cover_letter(letter_text: str, contact: dict, out_path: Path) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(11)

    if contact.get("name"):
        doc.add_paragraph(contact["name"])
    contact_line = " | ".join(filter(None, [contact.get("email"), contact.get("phone")]))
    if contact_line:
        doc.add_paragraph(contact_line)
    doc.add_paragraph("")

    for para in letter_text.split("\n\n"):
        doc.add_paragraph(para.strip())

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
